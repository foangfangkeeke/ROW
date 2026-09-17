from __future__ import annotations
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

def _rows(ws):
    values = list(ws.values)
    header = [str(x).strip() if x is not None else "" for x in values[0]]
    return [dict(zip(header, r)) for r in values[1:] if any(x is not None for x in r)]

@dataclass
class Data:
    path: Path; params: dict[str, Any]; nodes: list[int]; arcs: dict[int, dict]
    commodities: dict[int, dict]; periods: list[int]; scenarios: dict[int, dict]
    demand: dict[tuple[int,int], float]; cap: dict[tuple[int,int,int], tuple[float,float]]
    candidates: list[tuple[int,int]]; probabilities: dict[str, dict[int,float]]
    # In the revised benchmark, the second component of a candidate is a
    # coarse ROW block rather than a fine operating period.  Defaults retain
    # compatibility with the original R04 workbook and unit-test fixtures.
    blocks: list[int] = field(default_factory=list)
    block_of: dict[int,int] = field(default_factory=dict)
    candidate_info: dict[tuple[int,int],dict] = field(default_factory=dict)
    instance_id: str|None = None

    def row_cell(self, arc: int, fine_period: int) -> tuple[int,int]:
        return (arc, self.block_of[fine_period]) if self.block_of else (arc, fine_period)

    def activation_cost(self, cell: tuple[int,int]) -> float:
        return float(self.candidate_info.get(cell, {}).get('ROW_activation_cost', self.arcs[cell[0]].get('ROW_activation_cost_per_period', 0.0)))

    def resource(self, cell: tuple[int,int]) -> float:
        return float(self.candidate_info.get(cell, {}).get('resource_consumption', self.arcs[cell[0]].get('ROW_resource_weight', 0.0)))

def load_excel(path: str | Path, instance_id: str|None=None) -> Data:
    path = Path(path); wb = load_workbook(path, data_only=True, read_only=True)
    if 'Instance_Index' in wb.sheetnames and 'ROW_Candidates' in wb.sheetnames:
      return _load_block_benchmark(path, wb, instance_id)
    required = {"Parameters","Nodes","Physical_Arcs","Commodities","Time_Periods","Scenario_Index","Demand_Scenarios","Capacity_Base","Capacity_Factors"}
    missing = required-set(wb.sheetnames)
    if missing: raise ValueError(f"Missing required sheets: {sorted(missing)}")
    p = {r['parameter']: r['value'] for r in _rows(wb['Parameters'])}
    arcs = {int(r['arc_id']): r for r in _rows(wb['Physical_Arcs'])}
    com = {int(r['commodity_id']): r for r in _rows(wb['Commodities'])}
    periods = [int(r['period']) for r in _rows(wb['Time_Periods'])]
    scenarios = {int(r['scenario_id']):r for r in _rows(wb['Scenario_Index'])}
    dem = {(int(r['scenario_id']),int(r['commodity_id'])):float(r['realized_demand']) for r in _rows(wb['Demand_Scenarios'])}
    base = {(int(r['period']),int(r['arc_id'])):r for r in _rows(wb['Capacity_Base'])}
    factor = {(int(r['scenario_id']),int(r['period'])):r for r in _rows(wb['Capacity_Factors'])}
    cap={}
    for s in scenarios:
      for t in periods:
       for a,x in arcs.items():
        b=base[(t,a)]; f=factor[(s,t)]
        share=min(.25,max(.02,float(b['mean_background_share'])*float(b['background_exposure'])*float(f['capacity_factor'])))
        total=float(b['base_capacity_per_period']); row=total*float(p['ROW_capacity_share'])*float(p['autonomous_capacity_multiplier'])
        conventional=total*(1-share)-total*float(p['ROW_capacity_share'])
        cap[(s,a,t)]=(row, conventional)
    # A ROW decision exists for every physical candidate cell.  Cells late in
    # the horizon may not admit a new traversal, but can protect a traversal
    # that entered earlier; excluding them breaks the occupied-period coupling.
    cand=[(a,t) for a,x in arcs.items() if int(x['ROW_candidate']) for t in periods]
    probs={key:{s:float(r.get(f'probability_{key}',0) or 0) for s,r in scenarios.items() if int(r.get(f'in_{key}',0) or 0)} for key in ('S10','S20','S30')}
    d=Data(path,p,[int(r['node_id']) for r in _rows(wb['Nodes'])],arcs,com,periods,scenarios,dem,cap,cand,probs)
    audit(d); return d

def _load_block_benchmark(path: Path, wb, instance_id: str|None) -> Data:
    params={r['parameter']:r['value'] for r in _rows(wb['Parameters'])}
    index=_rows(wb['Instance_Index'])
    chosen=next((r for r in index if str(r['instance_id'])==str(instance_id)), None) if instance_id else next((r for r in index if '_S30' in str(r['instance_id'])), None)
    if chosen is None: raise ValueError(f'Unknown benchmark instance_id {instance_id!r}')
    variant=int(chosen['Y_variant'])
    # The original block workbook stores one row per (arc, Y variant).  The
    # consolidated R04 package instead has one arc row with Y-specific columns.
    raw_arcs=_rows(wb['Physical_Arcs'])
    consolidated='Y_variant' not in raw_arcs[0]
    if consolidated:
      y=f'Y{variant}'
      arcs={int(r['arc_id']):{**r,
        'autonomous_travel_periods':r['tau_A_periods'],
        'conventional_travel_periods':r['tau_H_periods'],
        'autonomous_unit_cost':r['c_A'],
        'conventional_unit_cost':r['c_H'],
        'base_capacity_per_period':r.get(f'capacity10_{y}',r.get(f'capacity15_{y}')),
      } for r in raw_arcs}
    else:
      arcs={int(r['arc_id']):r for r in raw_arcs if int(r['Y_variant'])==variant}
    com={int(r['commodity_id']):r for r in _rows(wb['Commodities'])}
    time_sheet='Fine_Periods' if consolidated else 'Time_Periods'
    periods=[int(r['period']) for r in _rows(wb[time_sheet])]
    time_rows={int(r['period']):r for r in _rows(wb[time_sheet])}
    block_of={t:int(time_rows[t]['ROW_block']) for t in periods}; blocks=sorted(set(block_of.values()))
    scenario_rows={int(r['scenario_id']):r for r in _rows(wb['Scenario_Index'])}
    scenarios={s:r for s,r in scenario_rows.items()}
    demand={(int(r['scenario_id']),int(r['commodity_id'])):float(r['realized_demand']) for r in _rows(wb['Demand_Scenarios'])}
    base={(int(r['period']),int(r['arc_id'])):r for r in _rows(wb['Capacity_Base']) if consolidated or int(r['Y_variant'])==variant}
    if consolidated:
      link={(int(r['scenario_id']),int(r['arc_id'])):float(r['link_Z']) for r in _rows(wb['Capacity_Link_Shocks'])}
      time={(int(r['scenario_id']),int(r['period'])):float(r['correlated_time_Z']) for r in _rows(wb['Capacity_Time_Shocks'])}
      lo,hi=(float(x) for x in str(params['background_share_bounds']).split(','))
    else:
      factor={(int(r['scenario_id']),int(r['period'])):r for r in _rows(wb['Capacity_Factors'])}
      lo=float(params['background_share_min']); hi=float(params['background_share_max'])
    row_share=float(params['ROW_capacity_share']); auto_mult=float(params['autonomous_capacity_multiplier'])
    cap={}
    for s in scenarios:
      for t in periods:
       for a,x in arcs.items():
        b=base[t,a]
        if consolidated:
          share=min(hi,max(lo,float(b['mean_background_share'])*float(b['background_exposure'])+.04*link[s,a]+.03*time[s,t]))
          total=float(b.get(f'capacity10_Y{variant}',b.get(f'capacity15_Y{variant}')))
          # R_H is pre-activation conventional capacity; Recourse subtracts
          # Delta when the corresponding ROW block is selected.
          cap[s,a,t]=(float(b[f'U_A_Y{variant}']), max(0.,total*(1.-share)))
        else:
          f=factor[s,t]
          share=min(hi,max(lo,float(b['mean_background_share'])*float(b['background_exposure'])*float(f['combined_factor'])))
          total=float(b.get('base_capacity_per_period',b.get('base_capacity_per_15min')))
          cap[s,a,t]=(total*row_share*auto_mult, max(0.,total*(1.-share-row_share)))
    candidates_rows=[r for r in _rows(wb['ROW_Candidates']) if (consolidated or int(r['Y_variant'])==variant) and int(r['candidate'])]
    if consolidated:
      candidates_rows=[{**r,'ROW_activation_cost':r[f'g_Y{variant}']} for r in candidates_rows]
    info={(int(r['arc_id']),int(r['ROW_block'])):r for r in candidates_rows}
    cand=sorted(info)
    probs={key:{s:float(r[f'probability_{key}']) for s,r in scenarios.items() if int(r[f'in_{key}'])} for key in ('S10','S20','S30')}
    budget_key={'Y1':'resource_budget_baseline','Y3':'resource_budget_baseline','Y5':'resource_budget_baseline'}
    p=dict(params); p['resource_budget']=float(chosen['baseline_resource_budget']) if consolidated else float(params[budget_key[f'Y{variant}']]); p['ROW_capacity_share']=row_share; p['autonomous_capacity_multiplier']=auto_mult
    d=Data(path,p,[int(r['node_id']) for r in _rows(wb['Nodes'])],arcs,com,periods,scenarios,demand,cap,cand,probs,blocks,block_of,info,str(chosen['instance_id']))
    audit(d); return d

def audit(d:Data):
    errs=[]
    if not d.candidates: errs.append('No candidate ROW cells')
    if d.periods != list(range(min(d.periods),max(d.periods)+1)): errs.append('Time_Periods must be consecutive integer periods')
    for a,x in d.arcs.items():
      ta,th=int(x['autonomous_travel_periods']),int(x['conventional_travel_periods'])
      if ta<1 or th<1: errs.append(f'arc {a}: travel periods must be >= 1')
      if ta > th: errs.append(f'arc {a}: tau_A > tau_H')
      for t in d.periods:
       for s in d.scenarios:
        _,rh=d.cap[(s,a,t)]
        if 'base_capacity_per_period' in x:
         delta=float(d.params['ROW_capacity_share'])*float(x['base_capacity_per_period'])
         if delta > rh+1e-8: errs.append(f'delta > R_H at scenario={s}, arc={a}, period={t}')
    for name in ('conventional_unit_cost','autonomous_unit_cost'):
      if any(float(x[name])<0 for x in d.arcs.values()): errs.append(f'negative cost/weight in {name}')
    for name,prob in d.probabilities.items():
      if prob and abs(sum(prob.values())-1.)>1e-8: errs.append(f'{name} probabilities sum to {sum(prob.values())}, not 1')
    if errs: raise ValueError('Data audit failed:\n'+'\n'.join(errs[:20]))

def audit_report(d:Data)->str:
 strict=sum(int(x['autonomous_travel_periods'])<int(x['conventional_travel_periods']) for x in d.arcs.values()); equal=len(d.arcs)-strict
 return '\n'.join([f'Excel: {d.path}',f'instance={d.instance_id}',f'nodes={len(d.nodes)}, physical_arcs={len(d.arcs)}, commodities={len(d.commodities)}, periods={len(d.periods)}',f'candidates={len(d.candidates)}, scenarios S10/S20/S30='+ '/'.join(str(len(d.probabilities[x])) for x in ('S10','S20','S30')),f'travel-time audit: tau_A<tau_H={strict}, tau_A=tau_H={equal}, tau_A>tau_H=0, zero-period arcs=0.',f'delta<=R_H verified for {len(d.scenarios)*len(d.arcs)*len(d.periods)} cells; nonnegative costs and scenario-probability normalization verified.', 'Mapping: Physical_Arcs(tail,head)->links; Fine_Periods.period->T; ROW_Candidates(arc_id,ROW_block)->C; Capacity_Base with link/time shocks -> (U_A,R_H); Demand_Scenarios -> d^omega; Scenario_Index -> p^omega.'])
